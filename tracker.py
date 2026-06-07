# tracker.py
# Creates and updates the Excel tracker file.
# Appends one row per invoice processed.

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os
from datetime import datetime


# Column headers for the tracker
HEADERS = [
    "Processed At",
    "Source File",
    "Invoice Number",
    "Vendor (Invoice)",
    "Vendor (PO)",
    "Invoice Date",
    "PO Number",
    "Invoice Amount",
    "PO Approved Amount",
    "Difference",
    "Status",
    "Discrepancy Note"
]

# Status colors
COLORS = {
    "MATCH":           "C6EFCE",   # green
    "AMOUNT_MISMATCH": "FFC7CE",   # red 
    "PO_NOT_FOUND":    "FFEB9C",   # yellow
    "AMOUNT_MISSING":  "FFEB9C",   # yellow
    "VENDOR_MISMATCH": "FFC7CE",   # red ow
}


def get_or_create_tracker(output_path):
    """
    Opens tracker.xlsx if it exists.
    Creates it with headers if it does not.
    Returns (workbook, worksheet).
    """
    if os.path.exists(output_path):
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Invoice Tracker"

        # Write header row
        for col_num, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font      = Font(bold=True, color="FFFFFF")
            #cell.fill      = PatternFill("solid", fgColor="2F5496")
            cell.alignment = Alignment(horizontal="center")

        # Set column widths
        widths = [18, 18, 16, 30, 30, 14, 12, 16, 18, 12, 18, 50]
        for col_num, width in enumerate(widths, start=1):
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col_num)
            ].width = width

    return wb, ws


def append_result(ws, result):
    """
    Appends one invoice result as a new row.
    Colors the row based on status.
    """
    invoice_amount = result.get("amount") or 0
    po_amount      = result.get("po_amount") or 0
    difference     = invoice_amount - po_amount if po_amount else None

    row_data = [
        datetime.now().strftime("%Y-%m-%d %H:%M"),
        result.get("source_file"),
        result.get("invoice_number"),
        result.get("vendor"),
        result.get("po_vendor"),
        result.get("date"),
        result.get("po_number"),
        invoice_amount,
        po_amount if po_amount else "N/A",
        difference if difference is not None else "N/A",
        result.get("status"),
        result.get("discrepancy_note")
    ]

    ws.append(row_data)

    # Color the entire row based on status
    status     = result.get("status", "")
    fill_color = COLORS.get(status, "FFFFFF")
    fill       = PatternFill("solid", fgColor=fill_color)

    for cell in ws[ws.max_row]:
        cell.fill = fill


def save_tracker(wb, output_path):
    """Saves the workbook to disk."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"  Tracker saved: {output_path}")